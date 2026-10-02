"""Export a reviewed report payload to Word; text checks are not semantic review.

This preview report contract is distinct from ledger.schema.json. Always preserve
the evidence ledger when available. Local Word sources use paragraph locators.
"""
from pathlib import Path
import argparse,hashlib,json,re
from collections import Counter
from read_manuscript import read_docx

def ledger_report(ledger, manuscript=None):
    """Render validated ledger facts without creating semantic review flags."""
    names={'direct':'直接支撑','partial':'部分支撑','insufficient':'相关但不足','not_found':'未找到','unable':'无法核准'}
    meta=ledger.get('report',{})
    sources={s['id']:s for s in ledger['sources']}
    claims=[]
    for c in ledger['claims']:
        evidence=[]
        for e in c['evidence']:
            item={'source_id':e['source_id'],'quote':e['quote'],'section':e['section']}
            if e.get('docx_location'):item['docx_location']=e['docx_location']
            else:
                item['pdf_page']=e['pdf_pages'][0];item['pdf_pages']=e['pdf_pages']
                item['print_page']='—'.join(str(p) for p in e['print_pages'] if p) or None
            evidence.append(item)
        location=c['target'].get('docx_location')
        claims.append({'id':c['id'],'heading':c.get('heading',c['target']['text'][:42]),
          'target_text':c['target']['text'],'section':c['target']['section'],
          'block_id':location['block_id'] if location else '主稿PDF第'+str(c['target']['pdf_page'])+'页',
          'status':names[c['status']],'reason':c['reason'],'recommendation':c['recommendation'],
          'source_ids':c['attributed_source_ids'],'evidence':evidence,'citation_role':c['citation_role'],
          'policy_checks':c['policy_checks']})
    return {'title':meta.get('title','引用表述来源核准报告'),
      'manuscript':{'title':meta.get('manuscript_title',(manuscript or {}).get('manuscript',{}).get('title','指定主稿'))},
      'sources':[dict(s,format='docx' if s['origin']['locator'].lower().endswith('.docx') else 'pdf') for s in sources.values()],
      'claims':claims,'highlights':meta.get('highlights',[]),'additional_notes':meta.get('additional_notes',[])}

def normalized(text,english=False):
    if english:
        text=re.sub(r'([A-Za-z])-\s*\n\s*([A-Za-z])',r'\1\2',text)
        return re.sub(r'\s+',' ',text).strip().replace('ﬁ','fi').replace('ﬂ','fl')
    return re.sub(r'\s+','',text).replace('','，').replace('\x00','')

def inside(root,locator):
    path=(Path(root)/locator).resolve()
    if not path.is_relative_to(Path(root).resolve()):raise ValueError('Source path escapes source_root')
    return path

def validate_report(data,corpus,manuscript,source_root,image_root=None):
    errors=[];seen=set();sources={s['id']:s for s in data['sources']}
    if len(sources)!=len(data['sources']):errors.append('Duplicate source ID')
    blocks={b['id']:b for b in manuscript['blocks']}
    segments={s['id']:{seg['id']:seg for p in s['pages'] for seg in p['segments']} for s in corpus['sources']}
    if manuscript['manuscript']['sha256']!=data['manuscript']['sha256']:errors.append('Manuscript digest mismatch')
    try:
        mainpath=inside(source_root,data['manuscript']['locator'])
        if hashlib.sha256(mainpath.read_bytes()).hexdigest()!=data['manuscript']['sha256']:errors.append('Manuscript changed')
    except (OSError,ValueError):errors.append('Manuscript unavailable or unsafe path')
    words={}
    for sid,s in sources.items():
        try:
            path=inside(source_root,s['origin']['locator'])
            if hashlib.sha256(path.read_bytes()).hexdigest()!=s['sha256']:errors.append('Source changed: '+sid)
            if path.suffix.lower()=='.docx':words[sid]={b['id']:b for b in read_docx(path,source_root)['blocks']}
        except (OSError,ValueError):errors.append('Source unavailable or unsafe path: '+sid)
    claim_ids=set()
    for claim in data['claims']:
        cid=claim['id']
        if cid in claim_ids:errors.append('Duplicate claim ID: '+cid)
        claim_ids.add(cid)
        if claim['target_text'] not in blocks.get(claim['block_id'],{}).get('text',''):errors.append('Target mismatch: '+cid)
        if claim['status'] not in ('直接支撑','部分支撑','相关但不足','未找到','无法核准'):errors.append('Invalid verdict: '+cid)
        if not claim.get('reason') or not claim.get('recommendation'):errors.append('Missing review: '+cid)
        if claim['status']=='未找到' and not claim.get('search_complete'):errors.append('Unfinished negative search: '+cid)
        if claim['status'] in ('直接支撑','部分支撑','相关但不足') and not claim.get('evidence'):errors.append('Missing evidence: '+cid)
        for q in claim.get('evidence',[]):
            key=(q['source_id'],q['quote'])
            seen.add(key)
            if q['source_id'] not in sources:errors.append('Unknown evidence source: '+cid);continue
            if q.get('region')!='body' or q.get('context_review')!='执行者复核':errors.append('Unreviewed body/context: '+cid)
            if q.get('docx_location'):
                block=words.get(q['source_id'],{}).get(q['docx_location']['block_id'],{})
                if block.get('text')!=q['quote'] or block.get('location')!=q['docx_location']['location']:errors.append('Word source quote/location mismatch: '+cid)
            elif q.get('spans'):
                try:
                    for s in q['spans']:
                        length=len(segments[q['source_id']][s['segment_id']]['text'])
                        if type(s['start']) is not int or type(s['end']) is not int or not 0 <= s['start'] < s['end'] <= length:
                            raise ValueError('Invalid evidence bounds')
                    raw=''.join(segments[q['source_id']][s['segment_id']]['text'][s['start']:s['end']] for s in q['spans'])
                    if q['quote'] not in (normalized(raw,True),normalized(raw,False)):errors.append('Evidence span mismatch: '+cid)
                except (KeyError,TypeError,ValueError):errors.append('Evidence span unavailable: '+cid)
            elif q.get('verification_image') and image_root:
                try:
                    for image in q['verification_image'].split(';'):
                        if not inside(image_root,image).is_file():errors.append('Verification image missing: '+cid)
                except ValueError:errors.append('Unsafe verification image path: '+cid)
            else:errors.append('Unbound quotation: '+cid)
    return {'errors':errors,'claims':len(data['claims']),'unique_quotes':len(seen),
            'semantic_review':'执行者复核，不是独立评审','image_check':'图像存在性不等于自动逐字核验',
            'ready_to_render':not errors}

def clean_title(title):
    title=re.sub(r'\s*\(z-library.*?\)','',title,flags=re.I)
    title=re.sub(r'^【.*?】','',title)
    title=title.replace('_',' ')
    return title.strip()

def export(data,output):
    from docx import Document
    from docx.shared import Pt,Cm,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_BREAK
    doc=Document();sec=doc.sections[0]
    sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.35);sec.right_margin=Cm(2.35)
    for sty in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[sty];style.font.name='宋体';style._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');style.font.color.rgb=RGBColor(0,0,0)
        for border in style._element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(5)
    for name,size in [('Title',20),('Heading 1',14),('Heading 2',12),('Heading 3',10.5)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
    sources={s['id']:s for s in data['sources']};used=list(dict.fromkeys(sid for c in data['claims'] for sid in c['source_ids']));labels={sid:f'R{i:02d}' for i,sid in enumerate(used,1)}
    counts=Counter(c['status'] for c in data['claims'])
    doc.add_paragraph(data['title'],style='Title')
    from datetime import date
    doc.add_paragraph('核对对象  '+data['manuscript']['title']+'  ｜  '+date.today().isoformat())
    doc.add_paragraph(f"本报告逐项核对所供主稿中的{len(data['claims'])}组引用，涉及{len(used)}种参考全文。判定为直接支撑{counts['直接支撑']}组、部分支撑{counts['部分支撑']}组、相关但不足{counts['相关但不足']}组、未找到{counts['未找到']}组、无法核准{counts['无法核准']}组。原始证据均选自正文；摘要、书目、版权页不作为论断证据，版权页仅用于版本核对。")
    doc.add_paragraph('直接支撑指原文对象与命题基本一致；部分支撑指需限定对象、范围、比较类别或推论强度；相关但不足指讨论有关主题而未支撑待核结论。判断针对本项节录，不自动覆盖同段其他命题。')
    doc.add_paragraph('定位采用Word段落编号及章节、参考PDF文件页码，并在可确认时并列原刊或原书页码。未确认的原页码不补猜。所供稿件为文献回顾或设计片段，本报告不代表全篇论文事实核查。')
    highlights=data.get('highlights',[])
    if highlights:
        doc.add_heading('优先处理的引用问题',1)
        for text in highlights:doc.add_paragraph(text,style='List Bullet')
    doc.add_heading('核准结果一览',1)
    table=doc.add_table(rows=1,cols=3);table.style='Table Grid';table.autofit=False
    table.columns[0].width=Cm(1.4);table.columns[1].width=Cm(10.3);table.columns[2].width=Cm(3.7)
    for cell,text in zip(table.rows[0].cells,['编号','待核命题与处理重点','判定']):cell.text=text
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for c in data['claims']:
        cells=table.add_row().cells;cells[0].text=c['id'];cells[1].text=c.get('heading',c['target_text'][:45]);cells[2].text=c['status']
    for cell in table.rows[0].cells:
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');cell._tc.get_or_add_tcPr().append(sh)
        for r in cell.paragraphs[0].runs:r.bold=True
    doc.add_heading('引用表述逐项核准',1)
    for c in data['claims']:
        doc.add_heading(c['id']+' '+c.get('heading','引用来源核准'),2)
        if c.get('citation_role'):
            doc.add_paragraph('引用类型  '+{'explicit':'明确引用','suggested':'建议补引','policy':'政策引用'}[c['citation_role']])
        doc.add_paragraph('主稿位置  '+c['section']+'  '+c['block_id'])
        p=doc.add_paragraph();p.add_run('主稿表述  ').bold=True;p.add_run(c['target_text'])
        p=doc.add_paragraph();r=p.add_run('判定  '+c['status']);r.bold=True
        if c['status']!='直接支撑':
            r.font.color.rgb=RGBColor.from_string('B00020')
            from docx.enum.text import WD_COLOR_INDEX
            r.font.highlight_color=WD_COLOR_INDEX.YELLOW
        for e in c['evidence']:
            p=doc.add_paragraph();p.add_run(labels[e['source_id']]+' '+clean_title(sources[e['source_id']]['title'])).bold=True
            p.paragraph_format.keep_with_next=True
            if e.get('docx_location'):
                loc='所供Word正文 '+e['docx_location']['block_id']+'  '+e['section'].split(' / ')[-1]+'  原书页码未核'
            else:
                pages=e.get('pdf_pages',[e['pdf_page']]);loc='正文 '+e['section']+'  PDF第'+ '—'.join(map(str,pages))+'页'
                loc+='  原页码'+str(e['print_page']) if e.get('print_page') else '  原页码未确认'
            p=doc.add_paragraph(loc);p.paragraph_format.keep_with_next=True
            p=doc.add_paragraph('“'+e['quote']+'”');p.paragraph_format.left_indent=Cm(.4);p.paragraph_format.right_indent=Cm(.25)
        p=doc.add_paragraph();p.add_run('支持边界  ').bold=True;p.add_run(c['reason'])
        p=doc.add_paragraph();p.add_run('处理建议  ').bold=True;p.add_run(c['recommendation'])
        if c.get('policy_checks'):
            policy=c['policy_checks'];levels={'formal':'正式原件','interpretation':'解读','secondary':'二手归纳'}
            doc.add_paragraph('政策核准层级  所称：'+levels[policy['claimed_level']]+'；实读：'+levels[policy['evidence_level']]+'。'+policy['version_note'])
    if data.get('additional_notes'):
        doc.add_heading('版本与相邻表述核准',1)
        for note in data['additional_notes']:doc.add_paragraph(note)
    doc.add_heading('来源清单与核准范围',1)
    for sid in used:
        s=sources[sid];format='所供Word转写全文' if s.get('format')=='docx' else '所供PDF全文'
        p=doc.add_paragraph(labels[sid]+'  '+clean_title(s['title'])+'  '+format)
        p.paragraph_format.space_after=Pt(3)
        for r in p.runs:r.font.size=Pt(9.5)
    doc.add_paragraph('判定来自本次执行者对所供材料的核对，尚无独立人工复审。缺少原件时只核准所供二手表述，不将其升级为政策原件、经典原著或不同版本的核准结果。')
    footer=sec.footer.paragraphs[0];footer.alignment=2
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('report_data',type=Path);p.add_argument('--corpus',type=Path);p.add_argument('--manuscript',type=Path)
    p.add_argument('--manuscript-root',type=Path)
    p.add_argument('--source-root',type=Path,required=True);p.add_argument('--image-root',type=Path);p.add_argument('--output',type=Path,required=True);p.add_argument('--qa-output',type=Path,required=True)
    a=p.parse_args();data=json.loads(a.report_data.read_text(encoding='utf8'))
    manuscript=json.loads(a.manuscript.read_text(encoding='utf8')) if a.manuscript else None
    if 'schema_version' in data:
        from validate_ledger import validate
        qa=validate(data,a.source_root,manuscript,a.manuscript_root or a.source_root,a.image_root)
        qa['errors']+=qa.get('blockers',[])
        if not qa['errors']:data=ledger_report(data,manuscript)
        qa['claims']=len(data['claims'])
    else:
        if not a.corpus or not manuscript:p.error('Legacy preview needs --corpus and --manuscript')
        qa=validate_report(data,json.loads(a.corpus.read_text(encoding='utf8')),manuscript,a.source_root,a.image_root)
    a.qa_output.parent.mkdir(parents=True,exist_ok=True);a.qa_output.write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf8')
    if qa['errors']:print(json.dumps(qa,ensure_ascii=False));return 2
    export(data,a.output);print(json.dumps({'output':str(a.output),'ready_to_render':True,'claims':qa['claims']},ensure_ascii=False));return 0

if __name__=='__main__':raise SystemExit(main())
