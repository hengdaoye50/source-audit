"""Synthetic regressions for defects exposed by the two local case runs."""
import copy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile

SCRIPTS=Path(__file__).resolve().parents[1]/'skills/source-audit/scripts'
sys.path.insert(0,str(SCRIPTS))
import discover_citations as discover
import export_report as reports
import read_manuscript as manuscript
import read_sources as sources
from test_manuscript import make_package

class CitationRegressionTests(unittest.TestCase):
    def test_note_anchor_is_a_candidate(self):
        with tempfile.TemporaryDirectory() as root:
            path=Path(root)/'main.docx';make_package(path)
            result=manuscript.read_docx(path,root)
        notes=[x for x in result['citation_candidates'] if x['kind']=='footnote_reference']
        self.assertEqual(len(notes),1)
        self.assertEqual(notes[0]['block_id'],'P000002')

    def test_note_inherits_anchor_section(self):
        with tempfile.TemporaryDirectory() as root:
            path=Path(root)/'main.docx';make_package(path)
            result=manuscript.read_docx(path,root)
        note=next(b for b in result['blocks'] if b['location'].get('note_id')=='2')
        self.assertEqual(note['section'],'一 研究评述')

    def test_omitted_author_is_retained(self):
        refs=discover.author_year('（甲，2022；2023）')
        self.assertEqual([(r['authors'],r['year']) for r in refs],[('甲','2022'),('甲','2023')])

    def test_multiple_english_notes_are_preserved(self):
        refs=discover.note_references('Alpha, "Synthetic First", 2014; Beta, "Synthetic Second", 2018')
        self.assertEqual([r['title'] for r in refs],['Synthetic First','Synthetic Second'])

    def test_multiple_versions_remain_candidates(self):
        files=[{'title':'甲 合成研究 2022 第一版','relative_path':'one.pdf'},
               {'title':'甲 合成研究 2022 再版','relative_path':'two.pdf'}]
        matches=discover.file_matches({'authors':'甲','year':'2022','title':'合成研究'},files)
        self.assertEqual(len(matches),2)
        self.assertTrue(all(m['candidate_only'] for m in matches))

    def test_null_text_is_rejected(self):
        self.assertFalse(sources.text_quality('正常候选文字'*30+'\x00')[0])

    def test_repeated_glyph_text_is_rejected(self):
        self.assertEqual(sources.text_quality('AAAA BBBB CCCC '+'normal words '*12)[1],'duplicated_glyph_runs')

    def test_encoded_glyph_text_is_rejected(self):
        self.assertFalse(sources.text_quality('正常候选文字'*30+'/C321')[0])

class ReportRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        main=self.root/'main.docx';make_package(main)
        self.manuscript=manuscript.read_docx(main,self.root)
        (self.root/'source.pdf').write_bytes(b'synthetic source only')
        self.corpus={'sources':[{'id':'S1','pages':[{'segments':[{'id':'p1.s1','text':'完整正文证据。'}]}]}]}
        self.data={'title':'自制核准报告','manuscript':self.manuscript['manuscript'],
          'sources':[{'id':'S1','title':'自制来源','sha256':hashlib.sha256(b'synthetic source only').hexdigest(),'origin':{'locator':'source.pdf'}}],
          'claims':[{'id':'G001','block_id':'P000002','section':'一 研究评述','target_text':'数字服务改善办事效率（甲，2024）。',
            'status':'直接支撑','reason':'完整支持自制命题','recommendation':'保留','source_ids':['S1'],
            'evidence':[{'source_id':'S1','quote':'完整正文证据。','section':'结论','pdf_page':1,'region':'body',
                'context_review':'执行者复核','spans':[{'segment_id':'p1.s1','start':0,'end':7}]}]}]}

    def tearDown(self):self.temp.cleanup()
    def validate(self):return reports.validate_report(self.data,self.corpus,self.manuscript,self.root)['errors']
    def test_bound_evidence_passes(self):self.assertEqual(self.validate(),[])
    def test_changed_source_is_rejected(self):
        (self.root/'source.pdf').write_bytes(b'changed');self.assertTrue(self.validate())
    def test_negative_span_is_rejected(self):
        self.data['claims'][0]['evidence'][0]['spans'][0]['start']=-7
        self.assertTrue(self.validate())
    def test_overflow_span_is_rejected(self):
        self.data['claims'][0]['evidence'][0]['spans'][0]['end']=999
        self.assertTrue(self.validate())
    def test_wrong_target_is_rejected(self):
        self.data['claims'][0]['target_text']='不在原稿中的命题';self.assertTrue(self.validate())
    def test_unreviewed_region_is_rejected(self):
        self.data['claims'][0]['evidence'][0]['region']='unknown';self.assertTrue(self.validate())
    def test_missing_evidence_is_rejected(self):
        self.data['claims'][0]['evidence']=[];self.assertTrue(self.validate())
    def test_unfinished_negative_is_rejected(self):
        self.data['claims'][0]['status']='未找到';self.data['claims'][0]['evidence']=[]
        self.assertTrue(self.validate())
    def test_escaping_source_is_rejected(self):
        self.data['sources'][0]['origin']['locator']='../outside.pdf';self.assertTrue(self.validate())
    def test_word_export_preserves_quote_and_repeating_header(self):
        from docx import Document
        path=self.root/'report.docx';reports.export(self.data,path);doc=Document(path)
        self.assertIn('“完整正文证据。”',[p.text for p in doc.paragraphs])
        self.assertTrue(doc.tables[0].rows[0]._tr.xpath('./w:trPr/w:tblHeader'))
        self.assertFalse(doc.styles['Title']._element.xpath('./w:pPr/w:pBdr'))

if __name__=='__main__':unittest.main()
