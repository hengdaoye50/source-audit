import copy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'skills/source-audit/scripts'))
from validate_ledger import validate
from export_report import ledger_report,export
from read_manuscript import read_docx

class UnifiedLedgerTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.data=json.loads((ROOT/'tests/fixtures/valid-ledger.json').read_text(encoding='utf8'))
        self.data['schema_version']='0.2'
    def tearDown(self):self.temp.cleanup()
    def word_source(self):
        from docx import Document
        path=self.root/'source.docx';doc=Document();doc.add_heading('研究结果',1)
        doc.add_paragraph('在本研究样本中，数字服务改善了办事效率，但未发现其显著增加就业。');doc.save(path)
        block=read_docx(path,self.root)['blocks'][1]
        source=self.data['sources'][0];source['pages']=[]
        source['origin']={'kind':'local','locator':'source.docx','snapshot_path':None}
        source['sha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        for c in self.data['claims']:
            for e in c['evidence']:
                e.update(source_sha256=source['sha256'],quote=block['text'],section=block['section'],
                    pdf_pages=[],print_pages=[],spans=[],docx_location={'block_id':block['id'],'location':block['location']})
        return block
    def test_existing_pdf_ledger_still_passes(self):
        self.assertTrue(validate(self.data)['ready_for_report'])
    def test_pdf_ledger_adapts_without_changing_verdict(self):
        data=ledger_report(self.data)
        self.assertEqual(data['claims'][0]['status'],'直接支撑')
        self.assertEqual(data['claims'][0]['evidence'][0]['quote'],self.data['claims'][0]['evidence'][0]['quote'])
    def test_word_source_passes_with_actual_paragraph(self):
        self.word_source();self.assertTrue(validate(self.data,self.root)['ready_for_report'])
    def test_word_source_requires_original_file(self):
        self.word_source();self.assertFalse(validate(self.data)['ready_for_report'])
    def test_word_source_wrong_locator_fails(self):
        self.word_source();self.data['claims'][0]['evidence'][0]['docx_location']['block_id']='P999999'
        self.assertFalse(validate(self.data,self.root)['ready_for_report'])
    def test_word_source_inferred_page_fails(self):
        self.word_source();self.data['claims'][0]['evidence'][0]['pdf_pages']=[1]
        self.assertFalse(validate(self.data,self.root)['ready_for_report'])
    def test_pdf_empty_binding_fails(self):
        self.data['claims'][0]['evidence'][0]['spans']=[]
        self.assertFalse(validate(self.data)['ready_for_report'])
    def test_transcription_requires_image_check(self):
        self.data['sources'][0]['pages'][0]['segments'][1]['verification_image']='page.png'
        self.assertFalse(validate(self.data)['ready_for_report'])
    def test_image_path_cannot_escape(self):
        self.data['sources'][0]['pages'][0]['segments'][1]['verification_image']='../outside.png'
        self.assertFalse(validate(self.data,image_root=self.root)['ready_for_report'])
    def test_policy_layer_reaches_report(self):
        c=self.data['claims'][0];c['citation_role']='policy'
        c['policy_checks']={'claimed_level':'formal','evidence_level':'secondary','quotation_verified':None,
          'dates_reviewed':True,'claims_originals_verified':False,'version_note':'只核准二手转述'}
        self.assertFalse(validate(self.data)['ready_for_report'])
        self.assertEqual(ledger_report(self.data)['claims'][0]['policy_checks'],c['policy_checks'])
    def test_export_preserves_partial_and_insufficient(self):
        for old,new in [('partial','部分支撑'),('insufficient','相关但不足'),('unable','无法核准')]:
            self.data['claims'][0]['status']=old
            self.assertEqual(ledger_report(self.data)['claims'][0]['status'],new)
    def test_export_preserves_word_location(self):
        from docx import Document
        block=self.word_source();data=ledger_report(self.data);path=self.root/'report.docx';export(data,path)
        text='\n'.join(p.text for p in Document(path).paragraphs)
        self.assertIn(block['id'],text);self.assertIn(block['text'],text)

if __name__=='__main__':unittest.main()
